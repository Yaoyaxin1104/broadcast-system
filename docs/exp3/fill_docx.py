# -*- coding: utf-8 -*-
"""填充实验3 Word 文档：广软广播站点歌与投稿系统"""
import os, copy, shutil
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = r"F:\实验3 数据库设计与数据库操作.docx"
TGT = r"C:\Users\yyxyz\IdeaProjects\broadcast-system\实验3 数据库设计与数据库操作（已完成）.docx"
ROOT = r"C:\Users\yyxyz\IdeaProjects\broadcast-system"
IMG = os.path.join(ROOT, "docs", "exp3")

shutil.copyfile(SRC, TGT)
doc = Document(TGT)
sec = doc.sections[0]
AVAIL = sec.page_width - sec.left_margin - sec.right_margin
AVAIL_CM = AVAIL / 360000
print("page avail cm:", AVAIL_CM)

# ---------------- helpers ----------------
def set_font(run, size=12, bold=False, code=False, eastasia="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    if code:
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        run._element.rPr.rFonts.set(qn("w:cs"), "Courier New")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), eastasia)
    else:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), eastasia)

def clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

def set_text(p, text, size=12, bold=False):
    clear_runs(p)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p

def shade(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def no_indent(p, line=240):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), str(line))
    sp.set(qn("w:lineRule"), "auto")

class Block:
    def __init__(self, cursor):
        self.cursor = cursor
    def _add_p(self):
        np = OxmlElement("w:p")
        self.cursor.addnext(np)
        self.cursor = np
        return Paragraph(np, doc.paragraphs[0]._parent)
    def body(self, text, bold=False, size=12):
        p = self._add_p()
        r = p.add_run(text)
        set_font(r, size=size, bold=bold)
        return p
    def label(self, text):
        return self.body(text, bold=True)
    def code(self, text, size=9):
        p = self._add_p()
        no_indent(p)
        shade(p)
        r = p.add_run(text if text else "")
        set_font(r, size=size, code=True)
        for t in r._element.findall(qn("w:t")):
            t.set(qn("xml:space"), "preserve")
        return p
    def code_block(self, s, size=9):
        for line in s.split("\n"):
            self.code(line if line else "", size=size)
    def caption(self, text):
        p = self._add_p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_indent(p)
        r = p.add_run(text)
        set_font(r, size=10.5, bold=False)
        return p
    def image(self, fname, max_h_cm=22.5):
        path = os.path.join(IMG, fname)
        iw, ih = Image.open(path).size
        w_cm = AVAIL_CM
        h_cm = w_cm * ih / iw
        if h_cm > max_h_cm:
            h_cm = max_h_cm
            w_cm = h_cm * iw / ih
        p = self._add_p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_indent(p)
        r = p.add_run()
        r.add_picture(path, width=Cm(w_cm))
        return p
    def table(self, rows, cols, widths_cm=None, header_shade=True):
        tb = doc.add_table(rows=rows, cols=cols)
        try:
            tb.style = "Table Grid"
        except Exception:
            pass
        self.cursor.addnext(tb._tbl)
        self.cursor = tb._tbl
        if widths_cm:
            for row in tb.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        return tb

def cell_text(cell, text, size=10.5, bold=False, center=False):
    p = cell.paragraphs[0]
    clear_runs(p)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 表格单元格内正文不允许首行缩进
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)

def zero_indent_table(tb):
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                ind = pPr.find(qn("w:ind"))
                if ind is None:
                    ind = OxmlElement("w:ind")
                    pPr.append(ind)
                ind.set(qn("w:firstLine"), "0")
                ind.set(qn("w:firstLineChars"), "0")

def find_p(sub):
    for p in doc.paragraphs:
        if sub in p.text:
            return p
    raise RuntimeError("not found: " + sub)

# ---------------- anchors ----------------
a_intro   = find_p("根据课程设计项目，整理业务需求")
b1 = find_p("一名用户可以下多笔订单")
b2 = find_p("一笔订单可以包含多条订单明细")
b3 = find_p("一门课程可以出现在多条订单明细中")
b4 = find_p("用户与课程之间通过")
a_output  = find_p("产出：需求文档、数据需求清单")
a_prompt2 = find_p("你是资深数据库架构师")
r1 = find_p("用户 1—N 订单")
r2 = find_p("订单 1—N 订单明细")
r3 = find_p("课程 1—N 订单明细")
shots = [p for p in doc.paragraphs if p.text.strip() == "截图："]
assert len(shots) == 3, len(shots)
a_prompt3 = find_p("请检查以下MySQL建表语句")
a_human   = find_p("人工修改：删除冗余字段")
a_ddlp    = find_p("请把上述E-R模型转换为MySQL 8关系模式")
a_exec    = find_p("在Navicat或命令行中执行DDL")
a_step21  = find_p("在MySQL中创建数据库，例如library_db")
a_yml     = find_p("在xxx-admin模块")
a_step23  = find_p("步骤3：创建持久层代码")
a_url     = find_p("启动项目后访问")
a_druid   = find_p("输入配置的监控用户名和密码")
a_task4   = find_p("统一响应数据格式、错误码枚举、自定义异常、全局通用异常")

# 表格按表头文字定位（插入新表后文档顺序会变化，不能用固定下标）
trel = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "表名")
tai  = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "实验环节")

# ---------------- 环境信息据实更新 ----------------
set_text(find_p("操作系统：Windows 10/11"), "操作系统：Windows 11")
set_text(find_p("JDK版本：JDK 18+"), "JDK版本：JDK 21（本次在 JDK 17 下构建验证通过）")
set_text(find_p("数据库：MySQL 8.0+"), "数据库：MySQL 8.0.43")
set_text(find_p("数据库客户端：Navicat 或 DBeaver"), "数据库客户端：MySQL 8.0 命令行（Navicat/DBeaver 亦可）")
set_text(find_p("AI工具：通义灵码"), "AI工具：豆包（Doubao）")

# ---------------- 任务一 步骤1 ----------------
set_text(a_intro, "根据课程设计项目「广软广播站点歌与投稿系统」，整理业务需求，形成需求文档与数据需求清单。系统面向校园广播站的日常业务，核心业务规则如下：")
rules = [
 "（1）系统有学生、广播站成员、教师三类用户，统一存储在 user 表中，通过 role 字段（student/staff/teacher）区分身份与权限。",
 "（2）一名学生可以提交多条点歌申请，每条点歌申请只属于一名学生（1:N）。",
 "（3）一名学生可以投稿多篇文章，每篇投稿只属于一名学生（1:N）。",
 "（4）一名广播站成员可以上传多条音频、编排多个节目单；每条音频、每个节目单只归属于一名成员（1:N）。",
 "（5）一名登录用户可以发表多条留言，每条留言只属于一名用户，广播站成员可对留言进行回复（1:N）。",
 "（6）点歌与投稿均需广播站成员审核，状态为 pending（待审核）/approved（通过）/rejected（驳回）；节目单有 draft（草稿）/published（已发布）两种状态。",
 "（7）经分析，系统中不存在 M:N 联系，无需拆分中间表；节目单中的播出歌单以快照文本（songs 字段）保存。",
]
set_text(b1, rules[0]); set_text(b2, rules[1]); set_text(b3, rules[2]); set_text(b4, rules[3])
bk = Block(b4._p)
for t in rules[4:]:
    bk.body(t)
# 数据需求清单
bk.label("数据需求清单（共 6 个数据实体）：")
data_req = [
 ("用户 user", "登录账号、密码（MD5密文）、真实姓名、学号、角色、电话、邮箱、账号状态、注册时间"),
 ("点歌单 song_request", "点歌学生、歌曲名、歌手、点歌留言、审核状态、审核时间、播放时间、申请时间"),
 ("投稿文章 article", "投稿学生、标题、正文、类型（news/story/poem）、审核状态、审核时间、投稿时间"),
 ("音频 audio", "上传成员、标题、文件路径、文件大小、时长、播放次数、上下架状态、上传时间"),
 ("留言 message", "留言用户、留言内容、回复内容、显示状态、留言时间"),
 ("节目单 program_schedule", "编排成员、节目日期、时间段、节目名称、节目内容、播出歌单快照、状态、发布时间"),
]
tb = bk.table(len(data_req) + 1, 2, widths_cm=[4.3, AVAIL_CM - 4.3])
cell_text(tb.rows[0].cells[0], "数据实体（表名）", bold=True, center=True)
cell_text(tb.rows[0].cells[1], "主要数据项", bold=True, center=True)
for i, (a, b) in enumerate(data_req, start=1):
    cell_text(tb.rows[i].cells[0], a)
    cell_text(tb.rows[i].cells[1], b)
zero_indent_table(tb)

# ---------------- 任务一 步骤2 ----------------
set_text(a_prompt2, "你是资深数据库架构师。请根据以下业务需求输出：1）实体清单及关键属性，标注主键候选；2）实体间关系及基数（1:1/1:N/M:N）；3）指出需要拆分为中间表的M:N关系。需求：广软广播站点歌与投稿系统，用户分学生、广播站成员、教师三类；学生在线点歌（歌名、歌手、留言）并投稿文章（新闻/故事/诗歌），成员审核点歌与投稿、上传节目音频、编排每日节目单（日期、时段、节目名、内容、歌单）；登录用户可给广播站留言，成员回复留言。")
set_text(r1, "用户（学生）1—N 点歌单（联系：提交）")
set_text(r2, "用户（学生）1—N 投稿文章（联系：撰写）")
set_text(r3, "用户（成员）1—N 音频（联系：上传）")
bk = Block(r3._p)
bk.body("用户 1—N 留言（联系：发表）")
bk.body("用户（成员）1—N 节目单（联系：编排）")
bk.body("AI 初稿输出 7 个实体（把“学生”“广播站成员”拆成两张表），并建议节目单与点歌单建立 M:N 中间表。人工定稿为 6 个实体（学生、成员、教师合并为 user，用 role 区分角色），各实体属性如下（带 * 为主键）：")
ents = [
 "user（id*、username、password、real_name、student_id、role、phone、email、status、create_time）",
 "song_request（id*、student_id、song_name、singer、message、status、audit_time、play_time、create_time）",
 "article（id*、student_id、title、content、type、status、audit_time、create_time）",
 "audio（id*、staff_id、title、file_path、file_size、duration、play_count、status、create_time）",
 "message（id*、user_id、content、reply、status、create_time）",
 "program_schedule（id*、staff_id、date、time_slot、program_name、content、songs、status、publish_time、create_time）",
]
for e in ents:
    bk.code(e, size=9)
bk.body("人工审查结论：① 学生与成员的登录、留言等属性完全一致，拆成两张表会导致登录与留言关系重复维护，故合并为 user 实体并用 role 区分；② 节目单 songs 字段保存的是播出时点歌单的快照，不需要随点歌记录联动，故不设中间表（属于情形B，已记入任务六 AI 使用声明）；③ 补全 AI 遗漏的留言回复 reply、音频播放次数 play_count 等关键属性。")
# 图1
bk = Block(shots[0]._p)
bk.image("er_chen.png")
bk.caption("图1  广软广播站点歌与投稿系统陈氏 E-R 图")

# ---------------- 任务一 步骤3 ----------------
set_text(a_prompt3, "请检查以下MySQL建表语句是否满足第三范式（3NF）。逐表说明：1）是否满足1NF；2）是否满足2NF；3）是否满足3NF；4）若存在冗余或更新异常，请指出具体字段并给出修改后的建表SQL。DDL：user、song_request、article、audio、message、program_schedule 六张表的 CREATE TABLE 语句（见本报告步骤4物理结构设计）。")
set_text(a_human, "人工修改：① AI 初稿在 song_request 中冗余了 student_name、student_no（学生姓名、学号）字段，违反 3NF，予以删除，学生信息通过 student_id 关联 user 表查询（在 SongRequestMapper.xml 中用 LEFT JOIN 实现）；② 采纳 AI 建议，为 username 增加唯一约束 uk_username，为高频过滤列 status、关联列 student_id/staff_id/user_id、日期列 date 增加普通索引；③ AI 建议节目单与点歌单拆 program_song 中间表，人工评估后认为 songs 是播出快照，不采纳，保留 TEXT 并补充注释（情形B）；④ AI 默认生成数据库级 FOREIGN KEY，人工定稿改为“逻辑外键+索引”，脚本末尾以注释保留等价物理外键语句；⑤ 密码字段按项目登录代码 UserServiceImpl 的实际实现改为 MD5 密文存储。")
# 关系模式表：6 行数据
rel_rows = [
 ("用户表 user", "id、username、password、real_name、student_id、role、phone、email、status、create_time", "id", "无"),
 ("点歌单表 song_request", "id、student_id、song_name、singer、message、status、audit_time、play_time、create_time", "id", "student_id（逻辑外键→user.id）"),
 ("投稿文章表 article", "id、student_id、title、content、type、status、audit_time、create_time", "id", "student_id（逻辑外键→user.id）"),
 ("音频表 audio", "id、staff_id、title、file_path、file_size、duration、play_count、status、create_time", "id", "staff_id（逻辑外键→user.id）"),
 ("留言表 message", "id、user_id、content、reply、status、create_time", "id", "user_id（逻辑外键→user.id）"),
 ("节目单表 program_schedule", "id、staff_id、date、time_slot、program_name、content、songs、status、publish_time、create_time", "id", "staff_id（逻辑外键→user.id）"),
]
while len(trel.rows) < len(rel_rows) + 1:
    tr = copy.deepcopy(trel.rows[-1]._tr)
    trel._tbl.append(tr)
for i, (a, b, c, d) in enumerate(rel_rows, start=1):
    cell_text(trel.rows[i].cells[0], a, size=10)
    cell_text(trel.rows[i].cells[1], b, size=9)
    cell_text(trel.rows[i].cells[2], c, size=10, center=True)
    cell_text(trel.rows[i].cells[3], d, size=9)
zero_indent_table(trel)
# 图2
bk = Block(shots[1]._p)
bk.image("ai_3nf_review.png", max_h_cm=22.0)
bk.caption("图2  AI 范式审查对话过程与逐表 1NF/2NF/3NF 结论")

# ---------------- 任务一 步骤4 ----------------
sql_text = open(os.path.join(ROOT, "sql", "broadcast_db.sql"), "r", encoding="utf-8").read()
bk = Block(a_ddlp._p)
bk.label("物理结构设计要点：")
phys = [
 "（1）存储引擎统一使用 InnoDB，字符集 utf8mb4、排序规则 utf8mb4_0900_ai_ci，完整支持中文及表情符号存储。",
 "（2）主键统一为 BIGINT AUTO_INCREMENT；create_time 设 DEFAULT CURRENT_TIMESTAMP 自动写入。",
 "（3）表间关联采用“逻辑外键+普通索引”：关联列建索引但不建数据库级 FOREIGN KEY，与 MyBatis-Plus 在应用层管理关联的方式一致，便于数据迁移与批量导入；脚本末尾以注释保留等价物理外键语句，需要强一致时可启用。",
 "（4）username 建唯一索引 uk_username 防止重复注册；status、student_id、staff_id、user_id、date 等高频过滤、关联、排序列建普通索引。",
 "（5）表名、字段名全部小写加下划线、单词用单数；每张表、每个字段均写 COMMENT 注释。",
]
for t in phys:
    bk.body(t)
bk.label("建表 DDL（完整脚本同时见项目 sql/broadcast_db.sql）：")
bk.code_block(sql_text, size=8)
bk = Block(a_exec._p)
bk.body("说明：为不破坏开发库 broadcast_db，以上脚本先在演示库 broadcast_db_demo 中完整重建并验证（图3~图5），验证通过后再用于项目库；图6为项目实际使用的 broadcast_db 库中的多表关联查询结果。")
bk = Block(shots[2]._p)
bk.image("ddl_execute.png"); bk.caption("图3  在命令行执行建库建表脚本（Query OK）并 SHOW TABLES 查看 6 张表")
bk.image("ddl_desc.png");    bk.caption("图4  DESC 查看 song_request 表结构、SHOW INDEX 查看主键与索引")
bk.image("test_data.png");   bk.caption("图5  演示库测试数据查询结果（用户、点歌单、投稿文章）")
bk.image("join_query.png");  bk.caption("图6  开发库 broadcast_db 中 LEFT JOIN 关联查询与按状态分组统计")

# ---------------- 任务二 ----------------
set_text(a_step21, "在 MySQL 8.0 中创建项目数据库 broadcast_db（字符集 utf8mb4、排序规则 utf8mb4_0900_ai_ci），并在命令行或 Navicat 中执行任务一生成的 sql/broadcast_db.sql，生成 user、song_request、article、audio、message、program_schedule 六张表并插入测试数据。建库语句：CREATE DATABASE IF NOT EXISTS broadcast_db DEFAULT CHARACTER SET utf8mb4 DEFAULT COLLATE utf8mb4_0900_ai_ci;")
yml_text = open(os.path.join(ROOT, "src", "main", "resources", "application-dev.yml"), "r", encoding="utf-8").read()
bk = Block(a_yml._p)
bk.body("项目为单模块工程，数据源配置位于 src/main/resources/application-dev.yml（application.properties 中通过 spring.profiles.active=dev 激活，服务端口 8089、上下文路径 /api）：")
bk.code_block(yml_text, size=8)
bk.body("配置说明：① spring.datasource.type 指定为 DruidDataSource，连接 broadcast_db 库；② 连接池 initial-size=5、min-idle=5、max-active=20，并配置 validation-query=SELECT 1 做连接保活；③ stat-view-servlet 开启 Druid 监控台，路径 /druid/*，账号 admin/admin123，reset-enable=false 禁止清空统计；④ web-stat-filter 采集 Web 请求与 JDBC 的关联统计；⑤ filter.stat 开启 SQL 统计与慢 SQL 日志（slow-sql-millis=1000），filter.wall 开启 SQL 防火墙；⑥ mybatis-plus 指定 Mapper XML 位置与实体别名包，开启下划线转驼峰 map-underscore-to-camel-case 和 SQL 标准输出日志。")

bk = Block(a_step23._p)
bk.body("（1）实体模型（PO）。entity 包下有 6 个实体类，与 6 张表一一对应，使用 MyBatis-Plus 的 @TableName、@TableId(type=IdType.AUTO) 注解和 Lombok 的 @Data；多表关联的扩展视图（如带学生姓名的点歌记录）由 Mapper XML 返回 LinkedHashMap（等价 VO），不新建冗余实体。以 User 为例：")
user_java = '''package com.grbroadcast.entity;

import com.baomidou.mybatisplus.annotation.IdType;
import com.baomidou.mybatisplus.annotation.TableId;
import com.baomidou.mybatisplus.annotation.TableName;
import lombok.Data;
import java.time.LocalDateTime;

@Data
@TableName("user")
public class User {
    @TableId(type = IdType.AUTO)
    private Long id;
    private String username;
    private String password;
    private String realName;
    private String studentId;
    private String role;
    private String phone;
    private String email;
    private Integer status;
    private LocalDateTime createTime;
}'''
bk.code_block(user_java, size=9)
bk.body("（2）DAO 层。dao 包下 6 个 Mapper 接口均继承 BaseMapper<T>，自动拥有单表增删改查能力；自定义 SQL 用 @Select 注解或 XML 实现。以 UserMapper 为例：")
user_mapper = '''package com.grbroadcast.dao;

import com.baomidou.mybatisplus.core.mapper.BaseMapper;
import com.grbroadcast.entity.User;
import org.apache.ibatis.annotations.Mapper;
import org.apache.ibatis.annotations.Select;

@Mapper
public interface UserMapper extends BaseMapper<User> {

    @Select("SELECT * FROM user WHERE username = #{username}")
    User findByUsername(String username);
}'''
bk.code_block(user_mapper, size=9)
bk.body("点歌模块的多表关联查询在 resources/mapper/SongRequestMapper.xml 中用 LEFT JOIN 实现，学生姓名、学号只在查询时关联取出，不冗余存储（符合 3NF）：")
xml_snip = '''<!-- 多表关联查询所有点歌记录（带学生信息） -->
<select id="getAllWithStudent" resultType="java.util.LinkedHashMap">
    SELECT
        sr.id, sr.student_id, sr.song_name, sr.singer,
        sr.message, sr.status, sr.audit_time, sr.play_time,
        sr.create_time,
        u.real_name as student_name,
        u.student_id as student_no,
        u.username
    FROM song_request sr
             LEFT JOIN user u ON sr.student_id = u.id
    ORDER BY sr.create_time DESC
</select>'''
bk.code_block(xml_snip, size=9)
bk.body("（3）分页插件。config 包下 MybatisPlusConfig 注册 MybatisPlusInterceptor 并加入 PaginationInnerInterceptor，使 BaseMapper 的 selectPage 分页生效：")
mp_cfg = '''package com.grbroadcast.config;

import com.baomidou.mybatisplus.annotation.DbType;
import com.baomidou.mybatisplus.extension.plugins.MybatisPlusInterceptor;
import com.baomidou.mybatisplus.extension.plugins.inner.PaginationInnerInterceptor;
import org.springframework.context.annotation.Bean;
import org.springframework.context.annotation.Configuration;

@Configuration
public class MybatisPlusConfig {

    @Bean
    public MybatisPlusInterceptor mybatisPlusInterceptor() {
        MybatisPlusInterceptor interceptor = new MybatisPlusInterceptor();
        interceptor.addInnerInterceptor(
                new PaginationInnerInterceptor(DbType.MYSQL));
        return interceptor;
    }
}'''
bk.code_block(mp_cfg, size=9)
bk.body("（4）数据流转：Controller 接收并校验请求 → Service（继承 ServiceImpl）组织业务逻辑 → Mapper（BaseMapper 单表 API 或 XML 关联 SQL）执行 SQL → Druid 连接池从连接池获取连接访问 MySQL → 查询结果封装为统一 Result 返回前端。")

# ---------------- 任务三 ----------------
set_text(a_url, "启动项目后访问：http://localhost:8089/api/druid（监控账号 admin / admin123，由 application-dev.yml 的 stat-view-servlet 配置）。")
bk = Block(a_druid._p)
bk.image("druid_login.png", max_h_cm=20.0);      bk.caption("图7  Druid 监控登录页")
bk.image("druid_index.png", max_h_cm=20.0);      bk.caption("图8  Druid 监控首页（版本、JVM、启动时间等）")
bk.image("druid_datasource.png", max_h_cm=20.0); bk.caption("图9  数据源监控（broadcast_db 连接信息、连接池与 StatFilter/WallFilter）")
bk.image("druid_sql.png", max_h_cm=20.0);        bk.caption("图10  SQL 监控（各 SQL 执行数、执行时间、读取行数）")
bk.image("druid_wall.png", max_h_cm=20.0);       bk.caption("图11  SQL 防火墙防御统计与表访问统计（非法次数为 0）")
bk.image("druid_weburi.png", max_h_cm=20.0);     bk.caption("图12  URI 监控（各 /api 接口请求数、JDBC 执行数、读取行数）")
bk.label("慢 SQL 排查与索引优化：")
bk.body("application-dev.yml 中 slow-sql-millis=1000、log-slow-sql=true 开启了慢 SQL 日志，SQL 监控页可按执行时间排序定位慢查询。以“查询待审核点歌” SELECT * FROM song_request WHERE status='pending' 为例：未建索引时 EXPLAIN 的 type=ALL（全表扫描）、possible_keys/key 均为 NULL、rows=4、filtered=25%；为 status 创建 idx_status 索引后，type=ref、key=idx_status、key_len=82、ref=const、rows=2、filtered=100%，扫描方式由全表扫描变为索引等值查询。演示数据量小，rows 绝对值差异不大，但执行计划已发生本质改变，数据量增大后优化效果显著。")
bk.image("explain_index.png"); bk.caption("图13  EXPLAIN 验证索引效果（上：无索引全表扫描；下：命中 idx_status）")
bk.body("另外，SQL 监控中歌名/歌手模糊查询为 LIKE '%关键词%'，前导通配符无法利用 BTree 索引，数据量增大后可改用 MySQL 全文索引；本实验保留 LIKE 实现，并要求该类查询尽量搭配 status 等可索引条件以缩小扫描范围。")

# ---------------- 任务四 ----------------
bk = Block(a_task4._p)
bk.body("项目在 common 包中设计了 4 个通用类，保证所有接口返回结构统一、错误码集中管理、异常处理与业务代码解耦：")
bk.body("（1）统一返回结果 Result：包含 code 状态码、msg 提示、data 数据三个字段，并提供 success/error 静态工厂方法：")
result_java = '''package com.grbroadcast.common;

import lombok.Data;

@Data
public class Result {
    private int code;
    private String msg;
    private Object data;
    public Result() {}

    public Result(int code, String msg, Object data) {
        this.code = code;
        this.msg = msg;
        this.data = data;
    }
    public static Result success(Object data) {
        return new Result(200, "success", data);
    }
    public static Result success(String msg, Object data) {
        return new Result(200, msg, data);
    }
    public static Result success(String msg) {
        return new Result(200, msg, null);
    }
    public static Result error(String msg) {
        return new Result(500, msg, null);
    }
    public static Result error(int code, String msg) {
        return new Result(code, msg, null);
    }
}'''
bk.code_block(result_java, size=9)
bk.body("（2）错误码枚举 ErrorCode，集中维护状态码与提示文案，避免魔法数字散落在代码中：")
bk.code_block('''public enum ErrorCode {

    SUCCESS(200, "操作成功"),
    PARAM_ERROR(400, "参数错误"),
    UNAUTHORIZED(401, "未登录或登录已过期"),
    FORBIDDEN(403, "无权限访问"),
    NOT_FOUND(404, "请求的资源不存在"),
    DUPLICATE(409, "数据已存在"),
    SYSTEM_ERROR(500, "系统繁忙，请稍后重试"),
    BUSINESS_ERROR(1001, "业务处理失败");

    private final int code;
    private final String msg;

    ErrorCode(int code, String msg) {
        this.code = code;
        this.msg = msg;
    }
    public int getCode() { return code; }
    public String getMsg() { return msg; }
}''', size=9)
bk.body("（3）自定义业务异常 BizException，在 Service 层按业务场景抛出，支持只传提示、传 ErrorCode 枚举或传自定义错误码三种构造方式：")
bk.code_block('''public class BizException extends RuntimeException {

    private final int code;

    public BizException(String message) {
        super(message);
        this.code = ErrorCode.BUSINESS_ERROR.getCode();
    }

    public BizException(ErrorCode errorCode) {
        super(errorCode.getMsg());
        this.code = errorCode.getCode();
    }

    public BizException(int code, String message) {
        super(message);
        this.code = code;
    }
    public int getCode() { return code; }
}''', size=9)
bk.body("（4）全局异常处理器 GlobalExceptionHandler，使用 @RestControllerAdvice 分层捕获业务异常、参数异常和兜底的 Exception，统一转换为 Result 返回：")
bk.code_block('''@Slf4j
@RestControllerAdvice
public class GlobalExceptionHandler {

    @ExceptionHandler(BizException.class)
    public Result handleBizException(BizException e) {
        log.warn("业务异常：code={}, msg={}", e.getCode(), e.getMessage());
        return Result.error(e.getCode(), e.getMessage());
    }

    @ExceptionHandler(IllegalArgumentException.class)
    public Result handleIllegalArgumentException(IllegalArgumentException e) {
        log.warn("参数异常：{}", e.getMessage());
        return Result.error(ErrorCode.PARAM_ERROR.getCode(), e.getMessage());
    }

    @ExceptionHandler(Exception.class)
    public Result handleException(Exception e) {
        log.error("系统异常", e);
        return Result.error(ErrorCode.SYSTEM_ERROR.getCode(),
                ErrorCode.SYSTEM_ERROR.getMsg());
    }
}''', size=9)
bk.body("调用链路：Service 校验不通过时抛出 throw new BizException(ErrorCode.NOT_FOUND) → @RestControllerAdvice 统一捕获 → Result.error(code, msg) 返回前端，前端按 code 统一提示；Controller 中无需编写 try-catch。")

# ---------------- 任务六 AI 声明表 ----------------
ai_rows = [
 ("任务一：概念结构设计（E-R 模型初稿）",
  "豆包",
  "“你是资深数据库架构师。请根据广软广播站点歌与投稿系统业务需求输出：1）实体清单及关键属性，标注主键候选；2）实体间关系及基数（1:1/1:N/M:N）；3）指出需要拆分为中间表的 M:N 关系。需求：学生点歌、投稿，成员审核、上传音频、编排节目单，用户留言。”",
  "输出 7 个实体（学生、成员分表）与 6 个联系，并建议节目单—点歌单拆 M:N 中间表",
  "合并学生/成员/教师为 user 实体并用 role 区分，定稿 6 个实体；否决节目单中间表，songs 改为快照文本；补全 reply、play_count 等遗漏属性",
  "三类用户登录属性一致，拆表会重复维护登录与留言关系；歌单是播出历史快照，无需与点歌记录联动（情形B）"),
 ("任务一：DDL 生成与 3NF 范式审查",
  "豆包",
  "“请把上述 E-R 模型转换为 MySQL 8 建表 DDL（InnoDB、utf8mb4，含主键、外键、必要索引、COMMENT），并逐表检查是否满足 1NF/2NF/3NF，指出冗余字段。”",
  "生成六表 DDL 与逐表范式结论，默认带数据库级 FOREIGN KEY，song_request 中冗余学生姓名、学号",
  "删除冗余姓名/学号，改由 LEFT JOIN 关联查询；物理外键改为逻辑外键+索引并在脚本末尾保留可选外键语句；补 uk_username 唯一键与 status、date 索引；密码改为 MD5 密文",
  "消除传递/部分依赖以满足 3NF，契合 MyBatis-Plus 应用层管理关联的实践；密码口径与项目 UserServiceImpl 的 MD5 登录逻辑保持一致"),
 ("任务三：慢 SQL 与索引优化",
  "豆包",
  "“song_request 按 status 过滤、按 create_time 排序，数据量大时如何优化？如何用 EXPLAIN 验证索引是否生效？”",
  "建议为 status 建立索引，用 EXPLAIN 的 type/key/rows/filtered 字段验证；提示 LIKE '%x%' 前导通配无法走 BTree 索引",
  "采纳建议建立 idx_status 并在演示库实测：type 由 ALL 变为 ref、key 命中 idx_status、rows 由 4 降为 2；规定模糊查询必须搭配 status 等可索引条件",
  "经 EXPLAIN 实测验证 AI 建议有效；模糊查询的使用限制是结合业务做的人工补充"),
]
for i, row in enumerate(ai_rows, start=1):
    for j, val in enumerate(row):
        cell_text(tai.rows[i].cells[j], val, size=9)
zero_indent_table(tai)

doc.save(TGT)
print("SAVED:", TGT)
