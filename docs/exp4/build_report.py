# -*- coding: utf-8 -*-
"""基于模板生成实验4报告"""
import os, shutil, fitz
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = r'F:\实验4 安全管理（认证与授权）实现.docx'
TGT = r'C:\Users\yyxyz\IdeaProjects\broadcast-system\实验4 安全管理（已完成）.docx'
SHOTS = r'C:\Users\yyxyz\IdeaProjects\broadcast-system\docs\exp4\shots'

shutil.copyfile(SRC, TGT)
doc = Document(TGT)
P = doc.paragraphs

def set_text(p, new):
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(new)

def _wrap(new_p, parent):
    return Paragraph(new_p, parent)

def new_par_after(p, text=''):
    np = OxmlElement('w:p')
    p._p.addnext(np)
    par = _wrap(np, p._parent)
    par.style = p.style
    if text:
        run = par.add_run(text)
        run.font.size = Pt(12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    return par

def new_par_before(p, text='', style=None):
    np = OxmlElement('w:p')
    p._p.addprevious(np)
    par = _wrap(np, p._parent)
    par.style = style or p.style
    if text:
        run = par.add_run(text)
        run.font.size = Pt(12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    return par

_fig = [0]
def img_after(p, name, caption, maxw=6.3, maxh=8.8):
    path = os.path.join(SHOTS, name + '.png')
    page = fitz.open(path)[0]
    ar = page.rect.height / page.rect.width
    w = maxw if ar * maxw <= maxh else maxh / ar
    np = new_par_after(p)
    np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    np.add_run().add_picture(path, width=Inches(w))
    _fig[0] += 1
    c = new_par_after(np, f'图{_fig[0]}  {caption}')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs:
        r.font.size = Pt(10.5)
    return c

def set_cell(cell, text):
    p0 = cell.paragraphs[0]
    set_text(p0, text)
    for extra in cell.paragraphs[1:]:
        set_text(extra, '')

# ---------- 实验环境 ----------
env = {
 28: '操作系统：Windows 11',
 29: '开发工具：IntelliJ IDEA 2025.1.2',
 30: 'JDK 版本：JDK 17（本机实际以 JDK 21 运行，向下兼容）',
 31: '构建工具：Maven 3.9.12',
 32: '版本管理：Git + Gitee 远程仓库',
 33: '数据库：MySQL 8.0',
 34: '缓存数据库：Redis 5.0.10（Windows 版，本次实验新增）',
 35: '数据库客户端：Navicat Premium 17；缓存客户端：redis-cli',
 36: '接口测试工具：Apifox（本机未安装，使用系统自带 curl 完成等效验证）',
 37: '核心依赖：Spring Boot 2.7.14、Spring Security 5.7.10、Spring Data Redis、MyBatis-Plus 3.5.3.1、Druid、MySQL Driver、Hutool 5.8.25、fastjson 1.2.83、Lombok',
}
for i, t in env.items():
    set_text(P[i], t)

# ---------- 实验准备 ----------
prep = {
 39: '课设项目「广软广播站点歌与投稿系统」为单模块 Spring Boot 工程，能够正常编译、启动（端口 8089，context-path 为 /api）。',
 40: '已创建数据库 broadcast_db，并完成业务表与初始数据的导入，学生、广播站成员等账号均已存在。',
 41: '已安装 Redis 5.0.10 并启动服务（默认端口 6379），能够使用 redis-cli 查看键与值。',
 42: '已在工程 pom.xml 中引入 Spring Security、Spring Data Redis、fastjson、Hutool 等依赖。',
 43: '已配置 Git 全局用户名与邮箱，能够在项目根目录提交代码（提交操作自行完成）。',
 44: '已能使用 curl 发送带请求头的 HTTP 请求，并能查看响应状态码与响应体。',
}
for i, t in prep.items():
    set_text(P[i], t)

# ---------- 任务一 步骤1 ----------
set_text(P[48], '按下表把依赖加入本工程 pom.xml；Spring Boot 启动器与 Redis 启动器版本由 Spring Boot 父工程统一管理，Hutool、fastjson 写明版本号。')
t_dep = doc.tables[1]
for row in t_dep.rows[1:]:
    set_cell(row.cells[2], '本工程')
a = img_after(P[49], 'code_pom', 'pom.xml 中新增的安全相关依赖')

# ---------- 步骤2 ----------
a = img_after(P[59], 'code_redis', 'RedisConfig 序列化配置')
a = img_after(a, 'runtime_redis', 'redis-cli 中查看到的缓存键与 JSON 值')

# ---------- 步骤3 ----------
a = img_after(P[70], 'code_jwt', 'JwtUtil 工具类')
a = img_after(a, 'runtime_jwt', '解开令牌后载荷仅含 userId、iat、exp')

# ---------- 步骤4 ----------
note = new_par_before(P[82], '说明：本系统角色固定为学生 student、广播站成员 staff、指导老师 teacher 三类，不设管理员账号，代码中保留超级管理员通配符分支；测试账号 staff1、student1，密码均为 123456，密码以 BCrypt 密文存储。')
a = img_after(P[82], 'code_loginuser', 'LoginUser 用户详情模型')
a = img_after(a, 'code_userdetails', 'UserDetailsServiceImpl 用户详情服务')

# ---------- 步骤5 ----------
set_text(P[96], '新建安全配置类，注册 SecurityFilterChain 类型的 Bean（Spring Security 5.7 起 WebSecurityConfigurerAdapter 已被弃用），并按表中逐项完成设置。')
adapt = new_par_before(P[98],
    '版本适配说明：本工程实际使用 Spring Boot 2.7.14、Spring Security 5.7.10。'
    '与模板基于的 Spring Boot 3.2 / Security 6 相比，URL 拦截规则使用 antMatchers'
    '（requestMatchers 的部分重载自 5.8 起才提供）；方法级安全使用 '
    '@EnableGlobalMethodSecurity(prePostEnabled = true) 开启；Security 5.7 已提供 '
    'AuthorizationManager 授权 API，任务二的自定义权限管理器同样按此实现；'
    'AuthenticationManager 通过 @Bean 暴露供登录接口调用。')
a = img_after(P[98], 'code_security', 'SecurityConfig 安全配置类')

# ---------- 步骤6 ----------
a = img_after(P[120], 'code_auth', 'AuthController 登录、当前用户与退出接口')
a = img_after(a, 'runtime_login', '调用登录接口成功返回令牌')

# ---------- 步骤7 ----------
a = img_after(P[134], 'code_filter', 'JwtAuthenticationTokenFilter 认证过滤器')
a = img_after(a, 'runtime_info', '携带令牌访问 /auth/info 成功返回')

# ---------- 步骤8 ----------
a = img_after(P[148], 'runtime_logout', '退出登录后原令牌立即失效返回 401')

# ---------- 新增步骤9：菜单 ----------
heading_style = P[47].style
body_style = P[48].style
h = new_par_before(P[165], '步骤9：设计菜单表、菜单树与角色菜单授权', heading_style)
b1 = new_par_before(P[165], '设计 menu 菜单表与 role_menu 角色菜单关联表。menu 表用 menu_type 字段区分目录 M、菜单 C、按钮 F 三类数据，权限标识按「模块:操作」命名，例如 song:audit、article:add；角色与菜单通过 role_menu 多对多关联。', body_style)
b2 = new_par_before(P[165], '菜单树采用两次遍历加映射表构建：第一次遍历把全部节点放入以菜单 ID 为键的映射表，第二次遍历为每个节点查找父节点并挂载、同时收集根节点。整体时间复杂度为 O(n)，既避免递归反复遍历列表，也避免层级过深时递归栈溢出。', body_style)
b3 = new_par_before(P[165], '角色菜单授权：角色列表接口返回各角色及其已选菜单；保存授权时在一个事务内先删除该角色的旧关联，再批量插入新关联，并自动补全所选菜单的父级目录，避免出现「能打开页面但上级目录不可见」。', body_style)
ip = new_par_before(P[165], style=body_style)
ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
ip.add_run().add_picture(os.path.join(SHOTS, 'code_menuservice.png'), width=Inches(6.3))
_fig[0] += 1
cap = new_par_before(P[165], f'图{_fig[0]}  MenuServiceImpl 菜单树构建与角色授权（先删后插、事务保护）', body_style)
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in cap.runs: r.font.size = Pt(10.5)
ip2 = new_par_before(P[165], style=body_style)
ip2.alignment = WD_ALIGN_PARAGRAPH.CENTER
ip2.add_run().add_picture(os.path.join(SHOTS, 'code_menucontroller.png'), width=Inches(6.3))
_fig[0] += 1
cap2 = new_par_before(P[165], f'图{_fig[0]}  MenuController 菜单与授权接口', body_style)
cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in cap2.runs: r.font.size = Pt(10.5)

# ---------- 任务二 步骤1 ----------
a = img_after(P[170], 'code_permservice', 'PermissionService 权限判断服务（Bean 名 ss）')

# ---------- 步骤2 ----------
a = img_after(P[188], 'code_permmanager', 'PermissionManager 自定义权限管理器')

# ---------- 步骤3 ----------
a = img_after(P[206], 'code_handlers', '401 与 403 两个异常处理器')
a = img_after(a, 'runtime_401', '无令牌访问受控接口返回 401')
a = img_after(a, 'runtime_403', '学生账号调用广播站成员接口返回 403')

# ---------- 步骤4：替换表 ----------
t_rep = doc.tables[4]
new_rows = [
 ('点歌模块接口', '方法内 Session/角色判断', 'song:list、song:add、song:audit、song:edit、song:delete'),
 ('投稿模块接口', '方法内 Session/角色判断', 'article:list、article:add、article:audit、article:edit、article:delete'),
 ('音频模块接口', '登录即可访问', 'audio:list、audio:upload、audio:delete'),
 ('留言模块接口', '登录即可访问', 'message:list、message:add、message:reply'),
 ('节目单模块接口', '方法内角色判断', 'program:list、program:publish、program:edit、program:delete'),
 ('菜单与授权接口', '本次实验新增', 'menu:list、menu:grant'),
]
data_rows = t_rep.rows[1:]
for row, vals in zip(data_rows, new_rows):
    for c, v in zip(row.cells, vals):
        set_cell(c, v)
# 删除多余的第7条数据行
extra = data_rows[len(new_rows)]
extra._tr.getparent().remove(extra._tr)

a = img_after(P[223], 'code_snip_song', 'SongController 方法级权限注解')
a = img_after(a, 'code_snip_article', 'ArticleController 方法级权限注解')
a = img_after(a, 'code_snip_audio', 'AudioController 方法级权限注解')
a = img_after(a, 'code_snip_message', 'MessageController 方法级权限注解')
a = img_after(a, 'code_snip_program', 'ProgramController 方法级权限注解')

# ---------- 删除任务三 Git 提交 ----------
body = doc.element.body
start_el = P[240]._p
els = list(body)
idx = els.index(start_el)
for el in els[idx:]:
    if el.tag == qn('w:sectPr'):
        continue
    body.remove(el)

doc.save(TGT)
print('SAVED', TGT)
